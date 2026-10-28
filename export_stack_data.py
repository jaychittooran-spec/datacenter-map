"""Export the STACK master workbook to validated website-ready JSON snapshots.

This deliberately uses stack_policy_* names; it does not overwrite the live site's
existing JSON contracts until the UI adapter has been reviewed.
"""
from pathlib import Path
from datetime import date, datetime
import json
from openpyxl import load_workbook
from refresh_stack_policy import KEYS, WEIGHTS, score_band

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'data'
OUT.mkdir(exist_ok=True)


def normalize(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, float) and (v != v or v in (float('inf'), -float('inf'))):
        return None
    return v


def records(sheet):
    items = iter(sheet.values)
    headers = [str(h).strip() if h is not None else '' for h in next(items)]
    output=[]
    for vals in items:
        if vals[0] is None or not str(vals[0]).strip():
            continue
        output.append({h:normalize(v) for h,v in zip(headers,vals) if h})
    return output


def write(filename, value):
    path=OUT/filename
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n',encoding='utf-8')
    print(f'EXPORTED: {path.relative_to(ROOT)}')


def main():
    workbook=ROOT/'STACK_Policy_Master.xlsx'
    if not workbook.exists():
        raise FileNotFoundError(f'Missing master: {workbook}')
    wb=load_workbook(workbook, read_only=True, data_only=True)
    expected={'State Scores','County Local','Policy Events'}
    if not expected.issubset(wb.sheetnames):
        raise ValueError(f'Missing worksheets: {sorted(expected-set(wb.sheetnames))}')
    states=records(wb['State Scores'])
    if len(states)!=50 or len({r['State'] for r in states})!=50:
        raise ValueError('Expected 50 unique states')
    for s in states:
        ratings=[float(s[key]) for key in KEYS]
        if any(v<1 or v>5 or (v*2)%1 for v in ratings):
            raise ValueError(f'Invalid dimension score: {s["State"]}')
        overall=round(20*sum(w*v for w,v in zip(WEIGHTS,ratings)),1)
        saved=s.get('Overall Score')
        if saved is not None and abs(overall-float(saved))>.051:
            raise ValueError(f'Weighted score mismatch: {s["State"]}')
        s['Overall Score']=overall
        s['Current Posture']=score_band(overall)
    counties=records(wb['County Local'])
    counties=[r for r in counties if r.get('County')]
    events=records(wb['Policy Events'])
    events=[r for r in events if r.get('Event Type') and r.get('Headline')]
    write('stack_policy_states.json',states)
    write('stack_policy_counties.json',counties)
    write('stack_policy_events.json',events)
    write('stack_policy_methodology.json',{'dimensions':[
        {'field':field,'weight':weight} for field,weight in zip(KEYS,WEIGHTS)],
        'formula':'20 × sum(dimension score × dimension weight)',
        'scale':'1 to 5 per dimension; overall 20 to 100',
        'current_policy_bands':{'Favorable':80,'Moderate':60,'Challenging':40,'Restrictive':0},
        'trend_and_forward_risk':'Reported separately from current policy score'})
    print(f'EXPORT COMPLETE: {len(states)} states; {len(counties)} counties; {len(events)} events')

if __name__=='__main__':
    main()
