"""Refresh the approved STACK presentation from the master workbook.

Only text, numeric score bars, and event rows are refreshed. The approved
slide layouts, logo, cover, footer, and corporate formatting are retained.

Run: python refresh_stack_policy.py
Dependencies: pip install python-pptx openpyxl
"""
from pathlib import Path
from datetime import date, datetime
from collections import defaultdict
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def source_file(name, downloaded_name):
    """Accept either GitHub filenames or the downloaded filenames."""
    for candidate in (ROOT / name, ROOT / downloaded_name):
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f'Missing {name}; place it next to this script.')


MASTER = source_file('STACK_Policy_Master.xlsx', 'STACK_Policy_Master_Download.xlsx')
TEMPLATE = source_file('STACK_Policy_Approved_Template.pptx', 'STACK_Policy_Approved_Template_Download.pptx')
OUTPUT = ROOT / 'STACK_Policy_Intelligence_Updated.pptx'
KEYS = ['Power & Energy (35%)', 'Water (15%)', 'Policy & Regulatory Environment (25%)', 'Community Sentiment (25%)']
WEIGHTS = (.35, .15, .25, .25)


def put(sh, value):
    """Replace content without replacing text-frame or font styling."""
    value = str(value if value is not None else '')
    p = sh.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = value
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.text = value
    for rest in sh.text_frame.paragraphs[1:]:
        for r in rest.runs:
            r.text = ''


def score_band(v):
    return 'Favorable' if v >= 80 else 'Moderate' if v >= 60 else 'Challenging' if v >= 40 else 'Restrictive'


def trend_text(v):
    return {'↑↑': 'Sharply Improving ↑↑', '↑': 'Improving ↑', '→': 'Stable →', '↓': 'Deteriorating ↓', '↓↓': 'Sharply Deteriorating ↓↓'}.get(str(v or '').strip(), str(v or ''))


def fmt_date(v):
    return v.strftime('%b %d, %Y') if isinstance(v, (date, datetime)) else str(v or '').strip()


def event_line(e, forward=False):
    when = str(e.get('Target Window') or '').strip() if forward else fmt_date(e.get('Event Date'))
    return '  |  '.join(filter(None, [when, str(e.get('State') or '').strip(), str(e.get('Headline') or '').strip()]))


def rows(sheet):
    data = list(sheet.values)
    return [dict(zip(data[0], r)) for r in data[1:] if r[0] is not None]


def add_event(slide, x, y, w, text):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(.47))
    tb.name = 'STACK_GENERATED_EVENT'
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = '• ' + text
    p.font.name = 'Aptos'
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(35, 55, 78)


def main():
    wb = load_workbook(MASTER, read_only=True, data_only=True)
    for name in ['State Scores', 'Policy Events', 'County Local']:
        if name not in wb:
            raise ValueError(f'Missing worksheet: {name}')
    states = rows(wb['State Scores'])
    if len(states) != 50 or len({s['State'] for s in states}) != 50:
        raise ValueError('Expected exactly 50 unique states')
    for s in states:
        vals = [float(s[k]) for k in KEYS]
        if any(v < 1 or v > 5 or (v * 2) % 1 for v in vals):
            raise ValueError(f'Invalid dimension score for {s["State"]}')
        expected = round(20 * sum(w * v for w, v in zip(WEIGHTS, vals)), 1)
        cached = s.get('Overall Score')
        if cached is not None and abs(expected - float(cached)) > .051:
            raise ValueError(f'Weighted score mismatch: {s["State"]}: {expected} vs {cached}')
        s['Overall Score'] = expected

    events = rows(wb['Policy Events'])
    recent = [e for e in events if 'current' in str(e.get('Event Type') or '').lower()]
    forward = [e for e in events if 'forward' in str(e.get('Event Type') or '').lower()]
    by_recent, by_forward = defaultdict(list), defaultdict(list)
    for e in recent:
        by_recent[str(e.get('State') or '').upper()].append(e)
    for e in forward:
        by_forward[str(e.get('State') or '').upper()].append(e)

    prs = Presentation(TEMPLATE)
    if len(prs.slides) != 59:
        raise ValueError('Template must have 59 slides')
    avg = sum(float(s['Overall Score']) for s in states) / 50
    kpi = {8: '50', 11: f'{avg:.1f}', 14: str(sum(float(s['Overall Score']) >= 80 for s in states)),
           17: str(sum(str(s.get('Forward Policy Risk') or '').upper() == 'HIGH' for s in states))}
    for i, v in kpi.items():
        put(prs.slides[1].shapes[i], v)

    s3 = prs.slides[2]
    group_defs = [('Favorable', 8, 9, 4), ('Moderate', 18, 19, 18), ('Challenging', 56, 57, 24), ('Restrictive', 106, 107, 4)]
    for label, header_idx, first_idx, capacity in group_defs:
        group = sorted([s for s in states if score_band(float(s['Overall Score'])) == label],
                       key=lambda s: (-float(s['Overall Score']), s['State']))
        if len(group) != capacity:
            raise ValueError(f'{label} group now has {len(group)} states, but approved slide has {capacity} slots. Update layout before rebuilding.')
        put(s3.shapes[header_idx], f'{label.upper()}  ({len(group)})')
        for j, state in enumerate(group):
            put(s3.shapes[first_idx + 2 * j], state['State'])
            put(s3.shapes[first_idx + 2 * j + 1], f"{float(state['Overall Score']):.1f}")

    by_state = {s['State'].upper(): s for s in states}
    for idx in range(4, 54):
        slide = prs.slides[idx]
        state_name = slide.shapes[1].text.strip().upper()
        if state_name not in by_state:
            raise ValueError(f'Unexpected state card: {state_name}')
        d = by_state[state_name]
        put(slide.shapes[8], d.get('Policy Thesis') or '')
        put(slide.shapes[10], f"{float(d['Overall Score']):.1f}")
        for text_idx, key, fill_idx in [(13, KEYS[0], 15), (18, KEYS[1], 20), (23, KEYS[2], 25), (28, KEYS[3], 30)]:
            value = float(d[key])
            put(slide.shapes[text_idx], f'{value:.1f}')
            slide.shapes[fill_idx].width = round(slide.shapes[fill_idx - 1].width * value / 5)
            slide.shapes[fill_idx].fill.fore_color.rgb = (RGBColor(56, 147, 70) if value >= 4 else
                                                          RGBColor(238, 177, 0) if value >= 3 else
                                                          RGBColor(230, 115, 20) if value >= 2 else RGBColor(202, 48, 42))
        put(slide.shapes[34], score_band(float(d['Overall Score'])))
        put(slide.shapes[36], trend_text(d.get('Trend')))
        put(slide.shapes[38], str(d.get('Forward Policy Risk') or '').capitalize())
        for shape in list(slide.shapes):
            if shape.name == 'STACK_GENERATED_EVENT' or (shape.has_text_frame and shape.text.startswith('• ')):
                shape._element.getparent().remove(shape._element)
        for n, e in enumerate(by_recent[state_name][:3]):
            add_event(slide, .95, 4.99 + n * .49, 6.45, str(e.get('Headline') or ''))
        for n, e in enumerate(by_forward[state_name][:3]):
            add_event(slide, 8.26, 4.99 + n * .49, 3.95, str(e.get('Headline') or ''))

    for slide_idx, items, is_forward in [(56, recent, False), (57, forward, True)]:
        slide = prs.slides[slide_idx]
        slots = [sh for sh in slide.shapes if sh.has_text_frame and '  |  ' in sh.text]
        if len(items) > len(slots):
            raise ValueError(f'Appendix slide {slide_idx + 1} has {len(slots)} slots for {len(items)} events')
        for j, sh in enumerate(slots):
            put(sh, event_line(items[j], is_forward) if j < len(items) else '')

    county_rows = rows(wb['County Local'])
    county_slots = [sh for sh in prs.slides[58].shapes if sh.has_text_frame and '  |  ' in sh.text]
    for j, sh in enumerate(county_slots):
        d = county_rows[j] if j < len(county_rows) else {}
        put(sh, '  |  '.join(filter(None, [str(d.get('State') or ''), str(d.get('County') or ''), str(d.get('County Tier') or '')])))
    if len(county_rows) > len(county_slots):
        print(f'NOTE: County appendix displays first {len(county_slots)} of {len(county_rows)} county rows; approved slide capacity is fixed.')
    prs.save(OUTPUT)
    with zipfile.ZipFile(OUTPUT) as z:
        if z.testzip():
            raise ValueError('PPTX archive integrity failed')
    check = Presentation(OUTPUT)
    if len(check.slides) != 59:
        raise ValueError('Unexpected slide count')
    print(f'BUILD COMPLETE: {OUTPUT.name} | 59 slides | 50 states | {len(recent)} current events | {len(forward)} forward catalysts')


if __name__ == '__main__':
    main()
